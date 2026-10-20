"""
Historial de accesos y cambios (auditoría) para administradores.

- admin: ve solo su colegio.
- super_admin: ve todo (o un colegio con ?school_id=).
- Cualquier otro rol: 403.

Las filas las escribe servicios/auditoria.py; aquí solo se consultan.
"""
import io
from datetime import datetime, timedelta

from flask import Blueprint, request, jsonify, send_file
from flask_jwt_extended import jwt_required
from sqlalchemy import func, or_

from models import db, AuditLog
from routes.access import current_claims, forbidden
from servicios.auditoria import EVENTOS

auditoria_bp = Blueprint('auditoria', __name__)

ETIQUETAS_EVENTO = {
    'login': 'Inicio de sesión', 'login_fallido': 'Inicio fallido', 'logout': 'Cierre de sesión',
    'crear': 'Creación', 'editar': 'Edición', 'eliminar': 'Eliminación',
    'descargar': 'Descarga', 'importar': 'Importación', 'accion': 'Acción',
}
ROLES = {'admin': 'Administrador', 'directivo': 'Directivo', 'profesor': 'Docente',
         'apoderado': 'Apoderado', 'alumno': 'Alumno', 'super_admin': 'Super admin'}

MAX_EXPORT = 50000


def _celda(v):
    """Evita que un texto («=HYPERLINK(...)») se interprete como fórmula en Excel."""
    if isinstance(v, str) and v[:1] in ('=', '+', '-', '@'):
        return "'" + v
    return v


def _alcance():
    """→ (query base, None) o (None, respuesta 403)."""
    c = current_claims()
    q = AuditLog.query
    if c['role'] == 'super_admin':
        sid = request.args.get('school_id', type=int)
        if sid:
            q = q.filter(AuditLog.school_id == sid)
        return q, None
    if c['role'] == 'admin' and c['school_id']:
        return q.filter(AuditLog.school_id == c['school_id']), None
    return None, forbidden('Solo un administrador puede ver el historial de accesos y cambios')


def _fecha(valor):
    try:
        return datetime.strptime(valor[:10], '%Y-%m-%d')
    except (TypeError, ValueError):
        return None


def _filtrar(q):
    """Aplica los filtros de la query string. Fechas en hora local del navegador
    (tz = minutos de getTimezoneOffset(); 0 = UTC)."""
    a = request.args
    tz = timedelta(minutes=a.get('tz', default=0, type=int) or 0)
    uid = a.get('user_id', type=int)
    if uid:
        q = q.filter(AuditLog.user_id == uid)
    evento = (a.get('event') or '').strip()
    if evento:
        q = q.filter(AuditLog.event.in_([e for e in evento.split(',') if e]))
    entidad = (a.get('entity') or '').strip()
    if entidad:
        q = q.filter(AuditLog.entity == entidad)
    desde = _fecha(a.get('desde'))
    if desde:
        q = q.filter(AuditLog.created_at >= desde + tz)
    hasta = _fecha(a.get('hasta'))
    if hasta:
        q = q.filter(AuditLog.created_at < hasta + timedelta(days=1) + tz)
    resultado = (a.get('resultado') or '').strip()
    if resultado == 'ok':
        q = q.filter(or_(AuditLog.status_code.is_(None), AuditLog.status_code < 400))
    elif resultado == 'error':
        q = q.filter(AuditLog.status_code >= 400)
    texto = (a.get('q') or '').strip()
    if texto:
        like = f'%{texto[:100]}%'
        q = q.filter(or_(AuditLog.description.ilike(like), AuditLog.user_name.ilike(like),
                         AuditLog.user_email.ilike(like), AuditLog.path.ilike(like)))
    return q


def _usuarios(base):
    """Usuarios distintos que aparecen en el historial (para el filtro)."""
    filas = (base.with_entities(AuditLog.user_id, func.max(AuditLog.user_name))
             .filter(AuditLog.user_id.isnot(None))
             .group_by(AuditLog.user_id).all())
    usuarios = [{'id': uid, 'nombre': nombre or f'Usuario #{uid}'} for uid, nombre in filas]
    return sorted(usuarios, key=lambda u: u['nombre'].lower())


@auditoria_bp.route('', methods=['GET'])
@auditoria_bp.route('/', methods=['GET'])
@jwt_required()
def listar():
    base, err = _alcance()
    if err:
        return err
    page = max(request.args.get('page', default=1, type=int) or 1, 1)
    per_page = min(max(request.args.get('per_page', default=50, type=int) or 50, 1), 200)
    q = _filtrar(base)
    total = q.order_by(None).count()
    filas = (q.order_by(AuditLog.created_at.desc(), AuditLog.id.desc())
             .offset((page - 1) * per_page).limit(per_page).all())
    entidades = sorted(e for (e,) in base.with_entities(AuditLog.entity)
                       .filter(AuditLog.entity.isnot(None)).distinct().all())
    return jsonify({
        'items': [f.to_dict() for f in filas],
        'total': total, 'page': page, 'per_page': per_page,
        'pages': max((total + per_page - 1) // per_page, 1),
        'usuarios': _usuarios(base),
        'entidades': entidades,
        'eventos': list(EVENTOS),
    }), 200


@auditoria_bp.route('/resumen', methods=['GET'])
@jwt_required()
def resumen():
    base, err = _alcance()
    if err:
        return err
    desde = datetime.utcnow() - timedelta(days=7)
    por_evento = dict(base.with_entities(AuditLog.event, func.count(AuditLog.id))
                      .filter(AuditLog.created_at >= desde)
                      .group_by(AuditLog.event).all())
    errores = (base.filter(AuditLog.created_at >= desde, AuditLog.status_code >= 400)
               .order_by(None).count())
    ultimos = (base.with_entities(AuditLog.user_id, func.max(AuditLog.user_name),
                                  func.max(AuditLog.role), func.max(AuditLog.created_at))
               .filter(AuditLog.event == 'login', AuditLog.user_id.isnot(None))
               .group_by(AuditLog.user_id).all())
    ultimos = sorted(ultimos, key=lambda r: r[3] or datetime.min, reverse=True)
    return jsonify({
        'desde': desde.isoformat() + 'Z',
        'por_evento': {e: por_evento.get(e, 0) for e in EVENTOS},
        'total_7d': sum(por_evento.values()),
        'errores_7d': errores,
        'ultimos_accesos': [{
            'user_id': uid, 'nombre': nombre or f'Usuario #{uid}', 'role': rol,
            'ultimo_login': (cuando.isoformat() + 'Z') if cuando else None,
        } for uid, nombre, rol, cuando in ultimos],
    }), 200


@auditoria_bp.route('/export.xlsx', methods=['GET'])
@jwt_required()
def exportar():
    base, err = _alcance()
    if err:
        return err
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    tz = timedelta(minutes=request.args.get('tz', default=0, type=int) or 0)
    filas = (_filtrar(base).order_by(AuditLog.created_at.desc(), AuditLog.id.desc())
             .limit(MAX_EXPORT).all())

    wb = Workbook()
    ws = wb.active
    ws.title = 'Auditoría'
    cab = ['Fecha y hora', 'Usuario', 'Email', 'Rol', 'Evento', 'Descripción', 'Entidad',
           'ID entidad', 'Resultado', 'Código', 'Método', 'Ruta', 'IP', 'Navegador']
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='334155')
        c.alignment = Alignment(vertical='center')
    for f in filas:
        ok = f.status_code is None or f.status_code < 400
        ws.append([_celda(v) for v in [
            (f.created_at - tz) if f.created_at else None,
            f.user_name or '', f.user_email or '', ROLES.get(f.role, f.role or ''),
            ETIQUETAS_EVENTO.get(f.event, f.event), f.description or '', f.entity or '',
            f.entity_id, 'OK' if ok else 'Error', f.status_code, f.method or '',
            f.path or '', f.ip or '', f.user_agent or '',
        ]])
    for celda in ws['A'][1:]:
        celda.number_format = 'dd-mm-yyyy hh:mm:ss'
    anchos = [19, 26, 30, 14, 17, 60, 22, 10, 10, 8, 8, 40, 16, 40]
    for i, w in enumerate(anchos):
        ws.column_dimensions[chr(ord('A') + i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nombre = f'auditoria_{datetime.utcnow():%Y%m%d_%H%M}.xlsx'
    return send_file(buf, as_attachment=True, download_name=nombre,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
