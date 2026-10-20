"""Historial de accesos y cambios (auditoría)."""
import io
import json
import uuid
from datetime import date, datetime, timedelta

import pytest

from conftest import CREDENTIALS, create_user, login
from models import db, AuditLog, Enrollment, School


def _ultimo(app, **filtros):
    with app.app_context():
        f = AuditLog.query.filter_by(**filtros).order_by(AuditLog.id.desc()).first()
        return f.to_dict() if f else None


def _count(app):
    with app.app_context():
        return AuditLog.query.count()


def _texto(fila):
    return json.dumps(fila, ensure_ascii=False, default=str)


@pytest.fixture(scope='module')
def otro_colegio(app):
    with app.app_context():
        s = School(name='Colegio Auditoría Otro', rut=f'77.{uuid.uuid4().int % 1000:03d}.222-2')
        db.session.add(s)
        db.session.commit()
        return s.id


# ── Inicio de sesión ─────────────────────────────────────────────────────────

def test_login_exitoso_queda_registrado_sin_clave(app, client, ids):
    email, clave = CREDENTIALS['admin']
    r = client.post('/api/auth/login', json={'email': email, 'password': clave},
                    headers={'X-Forwarded-For': '203.0.113.7, 10.0.0.1', 'User-Agent': 'pytest-ua'})
    assert r.status_code == 200
    fila = _ultimo(app, event='login', user_id=ids['admin'])
    assert fila['school_id'] == ids['school']
    assert fila['user_email'] == email
    assert fila['user_name'] == 'Carlos Administrador'
    assert fila['role'] == 'admin'
    assert fila['description'] == 'Inició sesión'
    assert fila['ip'] == '203.0.113.7'
    assert fila['user_agent'] == 'pytest-ua'
    assert fila['ok'] is True
    assert clave not in _texto(fila)


def test_login_fallido_guarda_email_y_colegio_sin_clave(app, client, ids):
    email = CREDENTIALS['profesor'][0]
    r = client.post('/api/auth/login', json={'email': email, 'password': 'clave-secreta-xyz'})
    assert r.status_code == 401
    fila = _ultimo(app, event='login_fallido', user_email=email)
    assert fila is not None
    assert fila['school_id'] == ids['school']
    assert fila['status_code'] == 401
    assert fila['ok'] is False
    assert fila['details']['motivo'] == 'Contraseña incorrecta'
    assert 'clave-secreta-xyz' not in _texto(fila)


def test_login_fallido_email_inexistente(app, client):
    r = client.post('/api/auth/login', json={'email': 'nadie@noexiste.cl', 'password': 'pw-123456'})
    assert r.status_code == 401
    fila = _ultimo(app, event='login_fallido', user_email='nadie@noexiste.cl')
    assert fila['school_id'] is None and fila['user_id'] is None
    assert 'pw-123456' not in _texto(fila)


def test_login_no_se_registra_dos_veces(app, client):
    antes = _count(app)
    login(client, *CREDENTIALS['directivo'])
    assert _count(app) == antes + 1


def test_logout_registrado(app, client, headers, ids):
    r = client.post('/api/auth/logout', headers=headers['directivo'])
    assert r.status_code == 200
    fila = _ultimo(app, event='logout', user_id=ids['directivo'])
    assert fila['description'] == 'Cerró sesión'


# ── Cambios ──────────────────────────────────────────────────────────────────

def test_request_que_modifica_queda_con_usuario_y_entidad(app, client, headers, ids, make_student):
    sid = make_student(first_name='Tomás', last_name='Auditado')
    r = client.put(f'/api/students/{sid}/profile', headers=headers['admin'],
                   json={'profile': {'comuna': 'Ñuñoa'}})
    assert r.status_code == 200, r.get_data(as_text=True)
    fila = _ultimo(app, endpoint='students.update_profile')
    assert fila['event'] == 'editar'
    assert fila['entity'] == 'Alumno'
    assert fila['entity_id'] == sid
    assert fila['user_id'] == ids['admin']
    assert fila['user_name'] == 'Carlos Administrador'
    assert fila['school_id'] == ids['school']
    assert fila['method'] == 'PUT'
    assert fila['status_code'] == 200
    assert fila['description'] == 'Editó la ficha de Tomás Auditado'
    assert fila['details']['campos'] == ['profile']


def test_retiro_describe_alumno_y_curso(app, client, headers, ids, make_student):
    sid = make_student(first_name='Rita', last_name='Retirada')
    with app.app_context():
        e = Enrollment(student_id=sid, course_id=ids['course'], year=2026, status='activo',
                       is_active=True, enrollment_date=date(2026, 3, 1))
        db.session.add(e)
        db.session.commit()
        eid = e.id
    r = client.delete(f'/api/matricula/enrollments/{eid}', headers=headers['admin'],
                      json={'motivo': 'Cambio de ciudad'})
    assert r.status_code == 200, r.get_data(as_text=True)
    fila = _ultimo(app, endpoint='matricula.retirar')
    assert fila['event'] == 'accion'
    assert fila['entity'] == 'Matrícula'
    assert fila['entity_id'] == eid
    assert fila['description'] == f"Retiró a Rita Retirada de {ids['course_name']}"


def test_accion_cerrar_y_rechazo_403_quedan_visibles(app, client, headers, ids):
    r = client.post('/api/anios/2031/cerrar', headers=headers['profesor'])
    assert r.status_code == 403
    fila = _ultimo(app, endpoint='anios.cerrar_anio')
    assert fila['event'] == 'accion'
    assert fila['entity'] == 'Año académico'
    assert fila['description'] == 'Cerró el año académico 2031'
    assert fila['status_code'] == 403 and fila['ok'] is False
    assert fila['user_id'] == ids['profesor']


def test_campos_sensibles_no_se_guardan(app, client, headers, ids):
    r = client.post('/api/auth/change-password', headers=headers['alumno'],
                    json={'current_password': 'clave-mala-777', 'new_password': 'nueva-888888'})
    assert r.status_code == 400
    fila = _ultimo(app, endpoint='auth.change_password')
    texto = _texto(fila)
    assert 'clave-mala-777' not in texto and 'nueva-888888' not in texto
    assert fila['details']['valores']['current_password'] == '[oculto]'
    assert fila['details']['campos'] == ['current_password', 'new_password']
    assert fila['details']['error']
    assert fila['description'] == 'Cambió su contraseña'


def test_saneamiento_de_valores():
    from servicios.auditoria import sanear
    png = 'data:image/png;base64,' + 'A' * 5000
    out = sanear({
        'signature_image': png, 'firma': 'x', 'access_token': 'tok', 'photoUrl': 'u',
        'archivo_base64': 'Zm9v', 'adjunto': 'abc',
        'observacion': png, 'blob_sin_nombre': 'B' * 300,
        'texto_largo': 'hola ' * 100, 'confirmar': True, 'profile': {'comuna': 'Ñuñoa'},
        'lista': list(range(50)),
    })
    for k in ('signature_image', 'firma', 'access_token', 'photoUrl', 'archivo_base64', 'adjunto'):
        assert out[k] == '[oculto]', k
    assert out['observacion'] == '[archivo omitido]'
    assert out['blob_sin_nombre'] == '[oculto]'
    assert len(out['texto_largo']) <= 201
    assert out['confirmar'] is True                 # «firma» no tapa «confirmar»
    assert out['profile'] == {'comuna': 'Ñuñoa'}    # «file» no tapa «profile»
    assert out['lista'] == '[50 elementos]'


def test_subida_de_archivo_guarda_solo_el_nombre(app, client, headers, make_student):
    sid = make_student()
    contenido = b'%PDF-1.4 contenido-secreto-del-pdf'
    r = client.post(f'/api/students/{sid}/documents', headers=headers['admin'],
                    data={'file': (io.BytesIO(contenido), 'certificado.pdf'), 'doc_type': 'otro'},
                    content_type='multipart/form-data')
    fila = _ultimo(app, endpoint='students.upload_document')
    assert fila is not None, r.status_code
    assert fila['details']['archivos'] == ['certificado.pdf']
    assert 'contenido-secreto' not in _texto(fila)


def test_get_normal_no_se_registra_y_descarga_si(app, client, headers):
    antes = _count(app)
    assert client.get('/api/anios', headers=headers['admin']).status_code == 200
    assert client.get('/api/auth/me', headers=headers['admin']).status_code == 200
    assert client.get('/api/auth/health').status_code == 200
    assert _count(app) == antes
    r = client.get('/api/matricula/plantilla.xlsx', headers=headers['admin'])
    assert r.status_code == 200
    assert _count(app) == antes + 1
    fila = _ultimo(app, endpoint='matricula.plantilla')
    assert fila['event'] == 'descargar'
    assert fila['description'] == 'Descargó la plantilla de matrícula'


def test_consultar_auditoria_no_se_registra(app, client, headers):
    antes = _count(app)
    client.get('/api/auditoria', headers=headers['admin'])
    client.get('/api/auditoria/export.xlsx', headers=headers['admin'])
    assert _count(app) == antes


# ── API de consulta ─────────────────────────────────────────────────────────

def _sembrar(app, school_id, marca, n, **extra):
    base = datetime.utcnow() - timedelta(days=1)
    with app.app_context():
        for i in range(n):
            db.session.add(AuditLog(
                school_id=school_id, user_id=extra.get('user_id'),
                user_name=extra.get('user_name'), role=extra.get('role'),
                event=extra.get('event', 'editar'), entity=extra.get('entity', 'Curso'),
                description=f'{marca} {i}', status_code=extra.get('status_code', 200),
                path='/api/courses/1', method='PUT', created_at=base - timedelta(minutes=i)))
        db.session.commit()


def test_admin_ve_solo_su_colegio(app, client, headers, ids, otro_colegio):
    _sembrar(app, otro_colegio, 'MARCA-OTRO', 3)
    _sembrar(app, ids['school'], 'MARCA-PROPIO', 2)
    r = client.get('/api/auditoria?q=MARCA-', headers=headers['admin'])
    assert r.status_code == 200
    descs = [i['description'] for i in r.get_json()['items']]
    assert len(descs) == 2 and all(d.startswith('MARCA-PROPIO') for d in descs)

    sa = login(client, 'superadmin@sistema.cl', 'super123')
    r = client.get('/api/auditoria?q=MARCA-', headers=sa)
    assert r.get_json()['total'] == 5
    r = client.get(f'/api/auditoria?q=MARCA-&school_id={otro_colegio}', headers=sa)
    assert r.get_json()['total'] == 3
    # Un admin no puede pedir otro colegio por parámetro
    r = client.get(f'/api/auditoria?q=MARCA-&school_id={otro_colegio}', headers=headers['admin'])
    assert r.get_json()['total'] == 2


@pytest.mark.parametrize('rol', ['profesor', 'directivo', 'apoderado', 'alumno'])
def test_otros_roles_403(client, headers, rol):
    for url in ('/api/auditoria', '/api/auditoria/resumen', '/api/auditoria/export.xlsx'):
        assert client.get(url, headers=headers[rol]).status_code == 403, url


def test_sin_token_401(client):
    assert client.get('/api/auditoria').status_code == 401


def test_filtros_y_paginacion(app, client, headers, ids):
    uid = create_user(app, role='profesor', first_name='Paula', last_name='Paginada')
    _sembrar(app, ids['school'], 'PAGINA-X', 7, user_id=uid, user_name='Paula Paginada',
             role='profesor')
    _sembrar(app, ids['school'], 'PAGINA-X borrado', 2, event='eliminar', entity='Nota',
             status_code=404)
    h = headers['admin']

    d = client.get('/api/auditoria?q=PAGINA-X&per_page=3', headers=h).get_json()
    assert d['total'] == 9 and d['pages'] == 3 and len(d['items']) == 3 and d['page'] == 1
    d3 = client.get('/api/auditoria?q=PAGINA-X&per_page=3&page=3', headers=h).get_json()
    assert len(d3['items']) == 3
    # Orden: más reciente primero
    fechas = [i['created_at'] for i in d['items']]
    assert fechas == sorted(fechas, reverse=True)

    d = client.get(f'/api/auditoria?q=PAGINA-X&user_id={uid}', headers=h).get_json()
    assert d['total'] == 7
    assert {'id': uid, 'nombre': 'Paula Paginada'} in d['usuarios']

    d = client.get('/api/auditoria?q=pagina-x&event=eliminar', headers=h).get_json()
    assert d['total'] == 2 and all(i['event'] == 'eliminar' for i in d['items'])
    d = client.get('/api/auditoria?q=PAGINA-X&entity=Nota', headers=h).get_json()
    assert d['total'] == 2
    d = client.get('/api/auditoria?q=PAGINA-X&resultado=error', headers=h).get_json()
    assert d['total'] == 2 and all(i['ok'] is False for i in d['items'])

    ayer = (datetime.utcnow() - timedelta(days=1)).date()
    d = client.get(f'/api/auditoria?q=PAGINA-X&desde={ayer + timedelta(days=2)}', headers=h).get_json()
    assert d['total'] == 0
    d = client.get(f'/api/auditoria?q=PAGINA-X&desde={ayer - timedelta(days=1)}&hasta={ayer + timedelta(days=1)}',
                   headers=h).get_json()
    assert d['total'] == 9
    d = client.get(f'/api/auditoria?q=PAGINA-X&hasta={ayer - timedelta(days=3)}', headers=h).get_json()
    assert d['total'] == 0

    # per_page con tope de 200
    d = client.get('/api/auditoria?per_page=5000', headers=h).get_json()
    assert d['per_page'] == 200


def test_resumen(app, client, headers, ids):
    login(client, *CREDENTIALS['admin'])
    d = client.get('/api/auditoria/resumen', headers=headers['admin']).get_json()
    assert d['por_evento']['login'] >= 1
    assert set(d['por_evento']) >= {'login', 'login_fallido', 'crear', 'editar', 'eliminar'}
    ultimos = {u['user_id']: u for u in d['ultimos_accesos']}
    assert ids['admin'] in ultimos and ultimos[ids['admin']]['ultimo_login']


def test_exportar_xlsx(app, client, headers, ids):
    _sembrar(app, ids['school'], 'EXPORTAR-Y', 4)
    _sembrar(app, ids['school'], '=HYPERLINK("http://x") EXPORTAR-Y', 1)
    r = client.get('/api/auditoria/export.xlsx?q=EXPORTAR-Y', headers=headers['admin'])
    assert r.status_code == 200
    assert r.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'attachment' in r.headers.get('Content-Disposition', '')
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(r.data)).active
    assert ws.cell(1, 1).value == 'Fecha y hora'
    assert ws.max_row == 6
    descs = [ws.cell(i, 6).value for i in range(2, 7)]
    assert sum(1 for x in descs if x.startswith('EXPORTAR-Y')) == 4
    assert any(x.startswith("'=") for x in descs)       # sin fórmulas inyectadas
